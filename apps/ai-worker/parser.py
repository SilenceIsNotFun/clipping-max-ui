import re

import pytesseract
from docx import Document
from PIL import Image
from pypdf import PdfReader

from schemas import ParsedDocument

URL_RE = re.compile(r"https?://[^\s)>\]]+")


def extract_links(text: str) -> list[str]:
    return list(dict.fromkeys(URL_RE.findall(text)))


def _parse_pdf(file_path: str) -> tuple[str, float]:
    reader = PdfReader(file_path)
    pages = [page.extract_text() or "" for page in reader.pages]
    text = "\n".join(pages)
    confidence = 0.9 if text.strip() else 0.1
    return text, confidence


def _parse_docx(file_path: str) -> tuple[str, float]:
    doc = Document(file_path)
    text = "\n".join(p.text for p in doc.paragraphs)
    confidence = 0.9 if text.strip() else 0.1
    return text, confidence


def _parse_image(file_path: str) -> tuple[str, float]:
    image = Image.open(file_path)
    data = pytesseract.image_to_data(image, output_type=pytesseract.Output.DICT)
    words = [w for w in data["text"] if w.strip()]
    text = " ".join(words)
    confidences = [int(c) for c in data["conf"] if c not in ("-1", -1)]
    confidence = (sum(confidences) / len(confidences) / 100) if confidences else 0.0
    return text, confidence


PARSERS = {
    "pdf": _parse_pdf,
    "docx": _parse_docx,
    "image": _parse_image,
}


def parse_document(file_path: str, doc_type: str) -> ParsedDocument:
    if doc_type not in PARSERS:
        raise ValueError(f"unsupported doc_type: {doc_type}")
    text, confidence = PARSERS[doc_type](file_path)
    return ParsedDocument(
        raw_text=text,
        extracted_links=extract_links(text),
        parsing_confidence=confidence,
    )
