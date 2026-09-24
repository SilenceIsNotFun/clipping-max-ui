# scripts/make_fixtures.py — run with: python scripts/make_fixtures.py
from docx import Document
from PIL import Image, ImageDraw
from reportlab.pdfgen import canvas

# sample.pdf
c = canvas.Canvas("apps/ai-worker/tests/fixtures/sample.pdf")
c.drawString(50, 750, "Reward Campaign Brief")
c.drawString(50, 730, "Source: https://example.com/brief-video")
c.save()

# sample.docx
doc = Document()
doc.add_paragraph("Reward Campaign Brief")
doc.add_paragraph("Source: https://example.com/brief-video")
doc.save("apps/ai-worker/tests/fixtures/sample.docx")

# sample.png
img = Image.new("RGB", (400, 100), color="white")
draw = ImageDraw.Draw(img)
draw.text((10, 40), "REWARD CAMPAIGN", fill="black")
img.save("apps/ai-worker/tests/fixtures/sample.png")
